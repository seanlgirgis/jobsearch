from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

root = Path(r"D:\Workarea\jobsearch")
reference = root / "resume_output" / "Sean_Girgis_Senior_Career_Resume.docx"
resume_out = root / "resume_output" / "Sean_Girgis_Fluidstack_IT_Infrastructure_Capacity_Analyst_Resume.docx"
letter_out = root / "resume_output" / "Sean_Girgis_Fluidstack_Cover_Letter.docx"

resume_content = [
    "Sean Girgis", None,
    "Senior Infrastructure Capacity and Data Engineer | Capacity Planning, Forecasting, Python, SQL, Automation",
    "Executive Profile",
    "Senior infrastructure capacity and data engineer with 20+ years of enterprise performance, telemetry, and operational analytics experience. Builds Python and SQL data pipelines that ingest, cleanse, reconcile, and retain infrastructure data for capacity planning, utilization analysis, and forecasting. Translates capacity findings into rightsizing, escalation, and proactive remediation decisions; modernizes recurring reporting and analyst workflows with AI-assisted automation.",
    "Core Competencies",
    "Capacity planning and infrastructure analytics | Demand and capacity forecasting | Python, Pandas, SQL, Oracle, MSSQL | Telemetry pipelines and data reconciliation | Utilization and bottleneck analysis | Capacity reporting and automated refresh | Performance engineering and observability | AI-assisted operational automation",
    "Professional Experience",
    "LTIMindtree | Capacity and Performance Specialist\t2026 - Present",
    "Apply AI-assisted automation to capacity reporting and recurring operational workflows, integrating data from JIRA, Microsoft SQL Server, and Splunk to improve repeatability, data availability, and analysis quality.",
    "Use Copilot, AI, and Python to organize, catalog, and support migration of Power BI reports, making reports easier to research and manage.",
    "Research an LLM-enabled team knowledge-base and analyst command-center concept supported by retrieval, Python, and PowerShell; the initiative is investigative and not represented as a production deployment.",
    "Citi | Senior Capacity and Data Engineer\tNov 2017 - Dec 2025",
    "Architected automated Python and Pandas ETL pipelines for P95 telemetry from 6,000+ endpoints, adding validation, cleansing, anomaly detection, and durable Oracle history for capacity analysis.",
    "Designed Oracle schemas and unified reporting to reconcile infrastructure telemetry, support seasonal utilization analysis, and surface capacity risks for planning decisions.",
    "Developed Prophet and scikit-learn forecasting models to identify infrastructure bottlenecks three to six months ahead and inform provisioning, consolidation, and rightsizing decisions.",
    "Managed enterprise capacity requests, monthly capacity reporting, peak-period analysis, and observability across BMC TrueSight, CA APM, and AppDynamics; translated analysis into operational action.",
    "G6 Hospitality | Performance Engineer\tMar 2017 - Nov 2017",
    "Managed Dynatrace AppMon and Synthetics for critical digital systems; led performance-data mining, bottleneck analysis, and AWS migration support for optimization recommendations.",
    "HCL / Entergy | APM Consultant\tJan 2017 - Mar 2017",
    "Supported CA Application Performance Management, Customer Experience Manager, and Application Delivery Analysis for utility systems.",
    "CA Technologies / TIAA CREF | Senior Consultant and CA APM Subject Matter Expert\t2011 - 2016",
    "Led enterprise CA APM implementations and upgrades across environments of roughly 4,000-6,000 agents and 50+ Enterprise Managers.",
    "Designed monitoring modules, dashboards, alerts, Golden Images, and Perl/KornShell extraction scripts; provided sizing guidance, troubleshooting, documentation, and client training.",
    "",
    "AT&T | Performance Test Engineer\tAug 2010 - Jul 2011",
    "Analyzed J2EE telecommunications applications under load, documented JDBC, thread, memory, CPU, and garbage-collection behavior, and automated monitoring with JMX and Wily Introscope.",
    "Sabre | Senior Systems and Data Migration Engineer\tMay 2008 - 2010",
    "Led migration of a high-throughput shopping engine from 200+ MySQL nodes to a six-node Oracle RAC cluster; optimized C++/OCCI processing and reduced hardware footprint by approximately 95% while maintaining sub-second latency.",
    "Earlier Experience | Software Architecture Development and Support\t1996 - 2008",
    "Built and supported high-availability C++ systems, billing interfaces, database and batch processes, and Linux/Unix operations across government, telecom, and enterprise environments.",
    "Selected Technical Projects",
    "HorizonScale (2025): Built a Python/PySpark capacity-forecasting project with Prophet, scikit-learn, and Streamlit; reduced forecasting cycles by approximately 90% and projected bottlenecks up to six months ahead.",
    "AI Job Search Pipeline (2026): Built a Python agentic workflow with FAISS semantic matching, LLM scoring and tailoring, document generation, quality gates, MCP integrations, and tracking across 100+ postings.",
    "Serverless Lakehouse: Designed an AWS S3, Glue, Athena, and Bedrock platform with a Text-to-SQL agent and Parquet/Snappy storage optimization.",
    "Technical Skills",
    "Capacity and infrastructure: Capacity planning, demand forecasting, capacity reporting, utilization analysis, bottleneck analysis, telemetry analysis, data quality and reconciliation, rightsizing, performance engineering",
    "Data and automation: Python, Pandas, SQL, Oracle, Microsoft SQL Server, ETL/ELT, PySpark/Spark, Airflow, Parquet, PowerShell, Linux/Unix",
    "Forecasting and platforms: Prophet, scikit-learn, trend and seasonal analysis, anomaly detection, Splunk, BMC TrueSight/TSCO, Dynatrace, CA APM, AppDynamics, AWS S3/Glue/Athena/Bedrock, Docker, Streamlit, AI-assisted automation",
    "Applied AI: AI-assisted automation, LLM workflows, RAG, FAISS, MCP, Copilot, Claude Code",
    "Education and Certifications",
    "Post-Graduate Diploma, Computer Engineering Technology / Computer Science - Humber College, Toronto | Bachelor of Science, Civil Engineering - Zagazig University, Egypt | Brainbench certifications: C++ Fundamentals, C++, C, Java 2.0, Object-Oriented Design, Mathematics",
]

resume = Document(reference)
assert len(resume.paragraphs) == len(resume_content)
for paragraph, text in zip(resume.paragraphs, resume_content):
    if text is None:  # Preserve the source contact line and its hyperlinks.
        continue
    for run in paragraph.runs:
        run.text = ""
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run.text = text
    if run.font.name:
        run._element.rPr.rFonts.set(qn('w:ascii'), run.font.name)
        run._element.rPr.rFonts.set(qn('w:hAnsi'), run.font.name)
resume.core_properties.title = "Sean Girgis Resume IT Infrastructure Capacity Analyst"
resume.core_properties.subject = "Fluidstack IT Infrastructure Capacity Analyst"
resume.core_properties.author = "Sean Girgis"
resume.save(resume_out)

letter = Document()
section = letter.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)
styles = letter.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(11)

name = letter.add_paragraph()
name.style = letter.styles['Title']
name.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = name.add_run('Sean Girgis')
run.font.name = 'Aptos Display'
run.font.size = Pt(20)
contact = letter.add_paragraph('Murphy, TX | 214-315-2190 | seanlgirgis@gmail.com | linkedin.com/in/sean-girgis-43bb1b5')
contact.paragraph_format.space_after = Pt(18)
for text in [
    'September 12, 2026',
    'Fluidstack Hiring Team',
    'Re: IT Infrastructure Capacity Analyst',
    'Dear Fluidstack Hiring Team,',
    'I am applying for the IT Infrastructure Capacity Analyst role. My career has centered on the same practical question your team is solving: what infrastructure will be needed, when will it be needed, where are the risks, and what action should follow from the data. I bring more than 20 years of enterprise capacity, performance, data, and operational analytics experience, including eight years as a Senior Capacity and Data Engineer at Citi.',
    'At Citi, I built Python and Pandas ETL pipelines for P95 telemetry from more than 6,000 endpoints, with validation, cleansing, anomaly checks, and Oracle history. I developed Prophet and scikit-learn models that identified capacity bottlenecks three to six months ahead, and I used capacity reporting, seasonal analysis, and observability data to support provisioning, consolidation, and rightsizing decisions. This work required reconciling imperfect operational data and turning the results into clear, timely actions for infrastructure stakeholders.',
    'In my current work, I am modernizing capacity operations through AI-assisted reporting and automated integration of JIRA, Microsoft SQL Server, and Splunk data sources. That work reinforces the habits that matter for this role: make capacity views repeatable, improve data quality, reduce manual refresh work, and surface risks early enough to act. My earlier performance-engineering and APM work also provides a strong understanding of the application, infrastructure, and telemetry behavior behind the planning numbers.',
    'Fluidstack\'s focus on rapidly scaling compute infrastructure is compelling. While I do not claim direct data-center supply-chain or ERP experience, I offer deep capacity-planning, forecasting, Python/SQL automation, telemetry analysis, and enterprise troubleshooting experience that I would bring to modeling demand, reconciling infrastructure gaps, and supporting proactive planning decisions.',
    'Thank you for your consideration. I would welcome the opportunity to discuss how my capacity-planning and infrastructure-analytics background can support Fluidstack\'s growth.',
    'Sincerely,',
    'Sean Girgis',
]:
    p = letter.add_paragraph(text)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.08
letter.core_properties.title = 'Sean Girgis Cover Letter Fluidstack'
letter.core_properties.subject = 'Fluidstack IT Infrastructure Capacity Analyst'
letter.core_properties.author = 'Sean Girgis'
letter.save(letter_out)

print(resume_out)
print(letter_out)
