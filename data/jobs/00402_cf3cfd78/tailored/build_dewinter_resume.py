from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

source = Path(r"D:\Workarea\jobsearch\resume_output\Sean_Girgis_Senior_Career_Resume.docx")
target = Path(r"D:\Workarea\jobsearch\resume_output\Sean_Girgis_DeWinter_FDE_Resume.docx")

content = [
    "Sean Girgis",
    None,
    "Senior Forward Deployed AI Applied AI Engineer | Python, Data, APIs, RAG, Agentic Workflows, Enterprise Automation",
    "Professional Summary",
    "Senior enterprise engineer with 20+ years of production systems, data, performance, and consulting experience. Translates operational and business problems into practical technical solutions using Python, APIs, data pipelines, automation, and applied AI workflows. Hands-on with RAG, agentic workflows, retrieval, prompt engineering, and AI-assisted development; experienced in stakeholder-facing delivery, enterprise troubleshooting, observability, reliability, and performance optimization.",
    "Core Competencies",
    "Applied AI and solution engineering | Python, data pipelines, and APIs | RAG, FAISS, agentic workflows, and MCP | Enterprise automation | Stakeholder discovery and technical problem framing | Reliability, troubleshooting, and performance engineering",
    "Professional Experience",
    "LTIMindtree | Capacity and Performance Specialist\t2026 - Present",
    "Translate recurring analyst and operational needs into AI-assisted capacity-reporting workflows that improve consistency, quality, and repeatability.",
    "Designed automated capacity-data workflows integrating JIRA, Microsoft SQL Server, and Splunk inputs to reduce fragmented manual collection and support repeatable analysis.",
    "Use Copilot, AI, and Python to organize, catalog, and support migration of Power BI reports, improving their discoverability and management.",
    "Research an LLM-enabled team knowledge-base and analyst command-center concept using retrieval, Python, and PowerShell; this work remains investigative rather than a production deployment.",
    "Citi | Senior Capacity and Data Engineer\tNov 2017 - Dec 2025",
    "Architected automated Python and Pandas ETL pipelines for P95 telemetry from 6,000+ endpoints, strengthening validation, cleansing, anomaly detection, and durable Oracle history.",
    "Designed Oracle schemas for historical retention and unified reporting, enabling seasonal analysis, executive dashboards, and capacity-risk forecasting.",
    "Developed Prophet and scikit-learn models to predict bottlenecks three to six months ahead and inform provisioning, consolidation, and rightsizing decisions.",
    "Partnered with operational stakeholders to troubleshoot production issues and deliver performance and reliability improvements across BMC TrueSight, CA APM, and AppDynamics.",
    "G6 Hospitality | Performance Engineer\tMar 2017 - Nov 2017",
    "Managed Dynatrace AppMon and Synthetics for critical digital systems; led performance-data mining, bottleneck analysis, and AWS migration support.",
    "HCL / Entergy | APM Consultant\tJan 2017 - Mar 2017",
    "Supported CA Application Performance Management, Customer Experience Manager, and Application Delivery Analysis for utility systems.",
    "CA Technologies / TIAA CREF | Senior Consultant and CA APM Subject Matter Expert\t2011 - 2016",
    "Led enterprise CA APM implementations and upgrades across environments of roughly 4,000-6,000 agents and 50+ Enterprise Managers.",
    "Designed monitoring modules, dashboards, alerts, Golden Images, and Perl/KornShell extraction scripts; provided sizing guidance, troubleshooting, documentation, and client training.",
    "",
    "AT&T | Performance Test Engineer\tAug 2010 - Jul 2011",
    "Analyzed J2EE telecommunications applications under load, documented JDBC, thread, memory, CPU, and garbage-collection behavior, and automated monitoring with JMX and Wily Introscope.",
    "Sabre | Senior Systems and Data Migration Engineer\tMay 2008 - 2010",
    "Led migration of a high-throughput shopping engine from 200+ MySQL nodes to a six-node Oracle RAC cluster; optimized C++/OCCI processing while maintaining sub-second latency.",
    "Earlier Experience | Software Architecture Development and Support\t1996 - 2008",
    "Built and supported high-availability C++ systems, billing interfaces, database and batch processes, and Linux/Unix operations across government, telecom, and enterprise environments.",
    "Selected Technical Projects",
    "AI Job Search Pipeline (2026): Built a Python agentic workflow with FAISS semantic matching, LLM scoring and tailoring, document generation, quality gates, MCP integrations, and tracking across 100+ postings.",
    "HorizonScale (2025): Built a Python/PySpark forecasting project with Prophet, scikit-learn, and Streamlit; reduced forecasting cycles by approximately 90% and projected bottlenecks up to six months ahead.",
    "Serverless Lakehouse: Designed an AWS S3, Glue, Athena, and Bedrock platform with a Text-to-SQL agent and Parquet/Snappy storage optimization.",
    "Technical Skills",
    "Applied AI: LLM applications, RAG, FAISS/vector search, agentic workflows, prompt engineering, MCP, AI evaluation, Copilot, Claude Code, AI-assisted development",
    "Software and data: Python, Pandas, SQL, Oracle, Microsoft SQL Server, APIs/integrations, ETL/ELT, PySpark/Spark, Airflow, Parquet, PowerShell, C++",
    "Cloud and reliability: AWS S3, Glue, Athena, Bedrock, Docker, Linux/Unix, Streamlit, Splunk, Dynatrace, CA APM, AppDynamics, BMC TrueSight, capacity planning, forecasting",
    "Education and Certifications",
    "Post-Graduate Diploma, Computer Engineering Technology / Computer Science - Humber College, Toronto | Bachelor of Science, Civil Engineering - Zagazig University, Egypt | Brainbench certifications: C++ Fundamentals, C++, C, Java 2.0, Object-Oriented Design, Mathematics",
]

doc = Document(source)
assert len(doc.paragraphs) == len(content), (len(doc.paragraphs), len(content))
for idx, (paragraph, text) in enumerate(zip(doc.paragraphs, content)):
    # The source contact line carries real hyperlink XML. Preserve it unmodified.
    if text is None:
        continue
    for run in paragraph.runs:
        run.text = ""
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run.text = text
    if run.font.name:
        run._element.rPr.rFonts.set(qn('w:ascii'), run.font.name)
        run._element.rPr.rFonts.set(qn('w:hAnsi'), run.font.name)
doc.core_properties.title = "Sean Girgis Resume Forward Deployed AI Engineer"
doc.core_properties.subject = "Resume for DeWinter Group Forward Deployed AI Engineer"
doc.core_properties.author = "Sean Girgis"
doc.save(target)
print(target)
