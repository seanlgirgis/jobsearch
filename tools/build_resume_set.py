from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "resume_output"
OUT.mkdir(exist_ok=True)

CONTACT = "Murphy, TX | 214-315-2190 | seanlgirgis@gmail.com | linkedin.com/in/sean-girgis-43bb1b5 | github.com/seanlgirgis"

def set_cell_margins(cell, top=60, start=80, bottom=60, end=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def hyperlink(paragraph, text, url):
    part = paragraph.part
    rel = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "1F4E79")
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single")
    rpr.extend([color, underline]); run.append(rpr)
    t = OxmlElement("w:t"); t.text = text; run.append(t); h.append(run)
    paragraph._p.append(h)

def keep_with_next(p):
    p.paragraph_format.keep_with_next = True

def configure(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(.48); sec.bottom_margin = Inches(.45)
    sec.left_margin = Inches(.58); sec.right_margin = Inches(.58)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"; normal.font.size = Pt(9.4); normal.font.color.rgb = RGBColor(30,30,30)
    normal.paragraph_format.space_after = Pt(1.8); normal.paragraph_format.line_spacing = 1.0
    for name, size, before, after in (("Title", 18, 0, 1), ("Subtitle", 10.5, 0, 3), ("Heading 1", 10.5, 5, 1), ("Heading 2", 9.7, 3, 0)):
        st = styles[name]
        st.font.name = "Aptos"; st.font.size = Pt(size); st.font.bold = name != "Subtitle"; st.font.color.rgb = RGBColor(0,0,0)
        st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after); st.paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Subtitle"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_ppr = styles["Title"].element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    if "Resume Contact" not in styles:
        s = styles.add_style("Resume Contact", WD_STYLE_TYPE.PARAGRAPH)
        s.font.name = "Aptos"; s.font.size = Pt(8.6); s.font.color.rgb = RGBColor(50,50,50)
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; s.paragraph_format.space_after = Pt(4)
    for s in styles:
        if hasattr(s, "font") and s.font:
            s.font.color.rgb = s.font.color.rgb or RGBColor(0,0,0)
    core = doc.core_properties
    core.author = ""; core.last_modified_by = ""; core.title = "Sean Girgis Resume"

def add_header(doc, headline, availability=None):
    p = doc.add_paragraph(style="Title"); p.add_run("Sean Girgis")
    p = doc.add_paragraph(style="Resume Contact")
    p.add_run("Murphy, TX | 214-315-2190 | seanlgirgis@gmail.com | ")
    hyperlink(p, "LinkedIn", "https://www.linkedin.com/in/sean-girgis-43bb1b5/")
    p.add_run(" | "); hyperlink(p, "GitHub", "https://github.com/seanlgirgis")
    p = doc.add_paragraph(style="Subtitle"); p.add_run(headline)
    if availability:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(availability); r.bold = True; r.font.size = Pt(9)

def heading(doc, text):
    return doc.add_paragraph(text, style="Heading 1")

def role(doc, company, title, dates):
    p = doc.add_paragraph(style="Heading 2")
    r = p.add_run(company); r.bold = True
    p.add_run(f" | {title}")
    tab = p.add_run("\t" + dates); tab.bold = True
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.28), WD_TAB_ALIGNMENT.RIGHT)
    keep_with_next(p)

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(.16); p.paragraph_format.first_line_indent = Inches(-.14)
    p.paragraph_format.space_after = Pt(1.2); p.paragraph_format.line_spacing = .96
    p.add_run(text)
    return p

def add_skills(doc, rows):
    for label, value in rows:
        p = doc.add_paragraph()
        r = p.add_run(label + ": "); r.bold = True
        p.add_run(value)

def save_doc(doc, path):
    doc.save(path)

def build_consulting():
    d = Document(); configure(d)
    add_header(d, "Python and AI Automation Consultant | Data Workflows | Reporting | Splunk and APM",
               "Available for select remote evening and weekend consulting engagements.")
    heading(d, "Consulting Profile")
    d.add_paragraph("Senior enterprise engineer who automates manual data, reporting, and operational workflows using Python and practical AI. Brings deep capacity planning, performance engineering, data integration, and observability experience to focused engagements that improve repeatability, quality, and decision speed.")
    heading(d, "Consulting Services")
    add_skills(d, [
        ("Workflow automation", "Python, PowerShell, APIs, AI-assisted development, repeatable multi-step workflows"),
        ("Data and reporting", "SQL, Pandas, ETL, data validation, capacity reporting, forecasting, report cataloging"),
        ("Observability", "Splunk, CA APM, Dynatrace, AppDynamics, BMC TrueSight, performance analysis"),
        ("Knowledge solutions", "RAG, FAISS, LLM workflows, searchable knowledge bases, technical mentoring"),
    ])
    heading(d, "Selected Outcomes and Projects")
    bullet(d, "Automated capacity-report preparation with AI and integrated data sources, standardizing the workflow, improving report quality, and reducing manual effort.")
    bullet(d, "Created automated capacity-data workflows using JIRA, Microsoft SQL Server, and Splunk inputs to reduce manual collection and make repeatable analysis possible.")
    bullet(d, "Applied Copilot, AI, and Python to organize, catalog, and support migration of a large Power BI report estate, improving searchability and management.")
    bullet(d, "Built an agentic job-search pipeline with Python, FAISS, sentence-transformers, LLM scoring and tailoring, document generation, quality gates, and application tracking across 100+ postings.")
    bullet(d, "Developed HorizonScale, a Python and PySpark capacity-forecasting project using Prophet and scikit-learn; reduced forecasting cycles by approximately 90% and projected bottlenecks up to six months ahead.")
    heading(d, "Professional Experience")
    role(d, "LTIMindtree", "Capacity and Performance Specialist", "2026 - Present")
    bullet(d, "Use practical AI and automation to improve capacity reporting, integrate operational data sources, and make recurring analyst workflows more consistent and efficient.")
    bullet(d, "Research a specialized team knowledge-base and analyst command-center concept using LLMs, retrieval, Python, and PowerShell; this work remains investigative rather than a production deployment.")
    role(d, "Citi", "Senior Capacity and Data Engineer", "Nov 2017 - Dec 2025")
    bullet(d, "Architected Python and Pandas ETL pipelines for P95 telemetry from 6,000+ endpoints, adding validation, cleansing, anomaly checks, and durable Oracle history.")
    bullet(d, "Developed Prophet and scikit-learn forecasting models to identify capacity risks three to six months ahead and support provisioning and rightsizing decisions.")
    bullet(d, "Managed capacity planning and observability across BMC TrueSight, CA APM, and AppDynamics; produced trend reporting, seasonal analysis, and executive dashboards.")
    role(d, "G6 Hospitality", "Performance Engineer", "Mar 2017 - Nov 2017")
    bullet(d, "Managed Dynatrace AppMon and Synthetics for critical digital systems, analyzed bottlenecks, supported an AWS migration, and led performance-data mining for optimization recommendations.")
    role(d, "Earlier Consulting and Engineering Experience", "APM, Performance, Data Migration and Software Engineering", "1996 - 2017")
    bullet(d, "Delivered enterprise APM implementations and upgrades for estates of roughly 4,000-6,000 agents; created monitoring modules, dashboards, alerts, extraction scripts, sizing guidance, and client training.")
    bullet(d, "Led migration of a high-throughput shopping engine from 200+ MySQL nodes to a six-node Oracle RAC cluster, reducing hardware footprint by approximately 95% while maintaining sub-second latency.")
    heading(d, "Technical Skills")
    d.add_paragraph("Python, SQL, Oracle, Microsoft SQL Server, Pandas, ETL/ELT, APIs, PowerShell, Linux/Unix, AWS S3/Glue/Athena/Bedrock, PySpark, Airflow, Docker, Parquet, Splunk, Power BI, CA APM, Dynatrace, AppDynamics, BMC TrueSight, Prophet, scikit-learn, Streamlit, RAG, FAISS, MCP, LLM and agentic workflows")
    heading(d, "Education and Certifications")
    d.add_paragraph("Post-Graduate Diploma, Computer Engineering Technology / Computer Science - Humber College, Toronto\nBachelor of Science, Civil Engineering - Zagazig University, Egypt\nBrainbench certifications: C++ Fundamentals, C++, C, Java 2.0, Object-Oriented Design, Mathematics")
    save_doc(d, OUT / "Sean_Girgis_Part_Time_Consulting_Resume.docx")

def build_senior():
    d = Document(); configure(d)
    add_header(d, "Senior Data Engineer | Capacity Planning Performance and AI-Driven Infrastructure Optimization")
    heading(d, "Executive Profile")
    d.add_paragraph("Senior data, capacity, and performance engineer with 20+ years delivering reliable data pipelines, forecasting, observability, and infrastructure analysis in regulated and high-scale enterprises. Combines Python, SQL, ETL, cloud data services, and application performance management with practical AI automation to improve operational decisions and reduce manual work.")
    heading(d, "Core Competencies")
    d.add_paragraph("Data Engineering and ETL | Capacity Planning and Forecasting | Performance Engineering | Python and SQL | AWS and PySpark | Data Quality | Observability and APM | AI-Assisted Automation | Technical Leadership")
    heading(d, "Professional Experience")
    role(d, "LTIMindtree", "Capacity and Performance Specialist", "2026 - Present")
    bullet(d, "Apply AI-assisted automation to capacity reporting and recurring operational workflows, integrating data from JIRA, Microsoft SQL Server, and Splunk to improve repeatability and report quality.")
    bullet(d, "Use Copilot, AI, and Python to organize, catalog, and support migration of a large Power BI report estate, making reports easier to research and manage.")
    bullet(d, "Research an LLM-enabled team knowledge-base and analyst command-center concept supported by retrieval, Python, and PowerShell; the initiative is investigative and not represented as production.")
    role(d, "Citi", "Senior Capacity and Data Engineer", "Nov 2017 - Dec 2025")
    bullet(d, "Architected automated Python and Pandas ETL pipelines for P95 telemetry from 6,000+ endpoints, replacing manual processes and strengthening validation, cleansing, and anomaly detection.")
    bullet(d, "Designed Oracle schemas for historical retention and unified reporting, enabling seasonal analysis, executive dashboards, and capacity-risk forecasting.")
    bullet(d, "Developed Prophet and scikit-learn models to predict bottlenecks three to six months ahead and inform provisioning, consolidation, and rightsizing decisions.")
    bullet(d, "Managed enterprise capacity requests, monthly reporting, peak-period analysis, and observability across BMC TrueSight, CA APM, and AppDynamics.")
    role(d, "G6 Hospitality", "Performance Engineer", "Mar 2017 - Nov 2017")
    bullet(d, "Managed Dynatrace AppMon and Synthetics for critical digital systems; led performance-data mining, bottleneck analysis, an upgrade from 6.5 to 7.0, TLS 1.2 enablement, and AWS migration support.")
    role(d, "HCL / Entergy", "APM Consultant", "Jan 2017 - Mar 2017")
    bullet(d, "Supported CA Application Performance Management, Customer Experience Manager, and Application Delivery Analysis for utility systems.")
    role(d, "CA Technologies / TIAA-CREF", "Senior Consultant and CA APM Subject Matter Expert", "2011 - 2016")
    bullet(d, "Led enterprise CA APM implementations and upgrades across environments of roughly 4,000-6,000 agents and 50+ Enterprise Managers.")
    bullet(d, "Designed monitoring modules, dashboards, alerts, Golden Images, and Perl/KornShell extraction scripts; provided sizing guidance, troubleshooting, documentation, and client training.")
    d.add_page_break()
    role(d, "AT&T", "Performance Test Engineer", "Aug 2010 - Jul 2011")
    bullet(d, "Analyzed J2EE telecommunications applications under load, documented JDBC, thread, memory, CPU, and garbage-collection behavior, and automated monitoring with JMX and Wily Introscope.")
    role(d, "Sabre", "Senior Systems and Data Migration Engineer", "May 2008 - 2010")
    bullet(d, "Led migration of a high-throughput shopping engine from 200+ MySQL nodes to a six-node Oracle RAC cluster; optimized C++/OCCI processing and reduced hardware footprint by approximately 95% while maintaining sub-second latency.")
    role(d, "Earlier Experience", "Software Architecture Development and Support", "1996 - 2008")
    bullet(d, "Built and supported high-availability C++ systems, billing interfaces, database and batch processes, and Linux/Unix operations across government, telecom, and enterprise environments.")
    heading(d, "Selected Technical Projects")
    bullet(d, "AI Job Search Pipeline (2026): Built a Python agentic workflow with FAISS semantic matching, LLM scoring and tailoring, document generation, quality gates, MCP integrations, and tracking across 100+ postings.")
    bullet(d, "HorizonScale (2025): Built a Python/PySpark forecasting project with Prophet, scikit-learn, and Streamlit; reduced forecasting cycles by approximately 90% and projected bottlenecks up to six months ahead.")
    bullet(d, "Serverless Lakehouse: Designed an AWS S3, Glue, Athena, and Bedrock platform with a text-to-SQL agent and Parquet/Snappy storage optimization.")
    heading(d, "Technical Skills")
    add_skills(d, [
        ("Data", "Python, SQL, Oracle, PL/SQL, Pandas, ETL/ELT, PySpark/Spark, Airflow, Parquet, dimensional modeling, data quality"),
        ("Cloud and platforms", "AWS S3, Glue, Athena, Bedrock; Docker; Linux/Unix; Streamlit"),
        ("Capacity and observability", "BMC TrueSight/TSCO, CA APM/Introscope, AppDynamics, Dynatrace, Splunk, Prophet, scikit-learn"),
        ("AI and automation", "LLM agents, RAG, FAISS, MCP, Copilot, AI-assisted development, PowerShell"),
    ])
    heading(d, "Education and Certifications")
    d.add_paragraph("Post-Graduate Diploma, Computer Engineering Technology / Computer Science - Humber College, Toronto | Bachelor of Science, Civil Engineering - Zagazig University, Egypt | Brainbench certifications in C++, C, Java, Object-Oriented Design, and Mathematics")
    save_doc(d, OUT / "Sean_Girgis_Senior_Career_Resume.docx")

if __name__ == "__main__":
    build_consulting(); build_senior()
