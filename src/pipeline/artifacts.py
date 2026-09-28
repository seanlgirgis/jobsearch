"""Evidence-locked intermediates and an adapter around the existing local renderers."""

from __future__ import annotations

import importlib.util
import re
import unicodedata
from pathlib import Path

from scripts.build_runtime_profile import (
    CREDENTIAL,
    FOUNDATION_WORDING,
    URL_OR_PATH,
    _strings,
    canonical,
    digest,
)
from src.pipeline.contracts import skill_inventory
from src.pipeline.storage import PipelineError, read_json, write_json, atomic_text


def public_text_check(value, email=""):
    text = "\n".join(_strings(value))
    normal = unicodedata.normalize("NFKC", text).casefold()
    compact = re.sub(r"[^a-z0-9]", "", normal)
    if any(x in compact for x in ("bankofamerica", "bofa", "bofamerica")) or re.search(
        r"\bboa\b", normal
    ):
        raise PipelineError("Prohibited client wording in model/artifact output.")
    if CREDENTIAL.search(text) or URL_OR_PATH.search(text.replace(email, "")):
        raise PipelineError("URL, path or credential pattern in model/artifact output.")
    if re.search(
        r"(?i)\b(?:CAPTAIN|lorem ipsum|your name|placeholder)\b|\[(?:company|role)\]",
        text,
    ):
        raise PipelineError(
            "Internal codename or placeholder in model/artifact output."
        )


def validate_package(package, profile, cover_requested):
    public_text_check(package)
    resume = package["resume"]
    if (package["cover"] is not None) != cover_requested:
        raise PipelineError(
            "Package cover selection does not match the requested mode."
        )
    prose = [resume["summary"]]
    if package["cover"]:
        cover = package["cover"]
        prose += [cover["intro"], *cover["body"], cover["conclusion"]]
    # Numbers and technical details remain in source-locked bullets, not free prose.
    free_prose = " ".join(prose)
    if re.search(r"\d", free_prose):
        raise PipelineError(
            "Free prose contains numerical claims; use source-locked evidence."
        )
    if FOUNDATION_WORDING.search(free_prose):
        raise PipelineError(
            "Free prose contains foundational-technology claims; use source-locked evidence."
        )
    if re.search(
        r"(?i)\b(?:certified|certification|PhD|doctorate|machine learning specialist|data science specialist)\b",
        " ".join(prose),
    ):
        raise PipelineError(
            "Unsupported credential or generic specialist claim in free prose."
        )
    if len(resume["summary"].split()) > 110:
        raise PipelineError("Resume summary exceeds the local word budget.")
    if package["cover"] and not 80 <= len(" ".join(prose[1:]).split()) <= 300:
        raise PipelineError("Cover must contain 80-300 words.")
    inventory = skill_inventory(profile)
    if not set(resume["skill_names"]) <= inventory.keys():
        raise PipelineError(
            "Selected skill is not in the allowed source-backed inventory."
        )
    evidence = {item["id"]: item for item in profile["evidence"]}
    selected_ids = []
    for section in ("experience", "projects"):
        for selection in resume[section]:
            selected_ids.append(selection["evidence_id"])
            item = evidence.get(selection["evidence_id"])
            if item is None or (item["kind"] == "project") != (section == "projects"):
                raise PipelineError("Invalid evidence section or ID.")
            if not set(selection["bullet_ids"]) <= set(item["bullet_ids"]):
                raise PipelineError("Bullet does not belong to its cited evidence.")
            if len(set(selection["bullet_ids"])) != len(selection["bullet_ids"]):
                raise PipelineError("Duplicate bullet selection.")
    if len(set(selected_ids)) != len(selected_ids) or "current" not in selected_ids:
        raise PipelineError(
            "Current work must be included once; evidence cannot be duplicated."
        )


def materialize(package, profile):
    """The model cannot supply contact details, job titles/dates, education or metrics."""
    selected = package["resume"]
    identity = profile["identity"]
    evidence = {x["id"]: x for x in profile["evidence"]}
    bank = {x["id"]: x["text"] for x in profile["bullet_bank"]}
    experience, projects = [], []
    for selection in selected["experience"]:
        item = evidence[selection["evidence_id"]]
        current = item["kind"] == "current"
        experience.append(
            {
                "company": item["employer"],
                # This is a presentation label, not an asserted official job title.
                # Current employment's formal title is intentionally unconfirmed.
                "title": (
                    current_display_label(item["function"])
                    if current
                    else item["title"]
                ),
                "start_date": "" if current else str(item["start"]),
                "end_date": "Present" if current else str(item["end"]),
                "bullets": [bank[x] for x in selection["bullet_ids"]],
            }
        )
    for selection in selected["projects"]:
        item = evidence[selection["evidence_id"]]
        projects.append(
            {
                "name": item["name"] + " (project)",
                "description": " ".join(bank[x] for x in selection["bullet_ids"]),
                "technologies": [],
            }
        )
    inventory = skill_inventory(profile)
    skills = {}
    for name in dict.fromkeys(selected["skill_names"]):
        group = inventory[name]
        skills[group] = ", ".join(filter(None, (skills.get(group), name)))
    resume = {
        "personal": {
            "full_name": identity["name"],
            "preferred_title": "",
            "phone": identity["phone"],
            "email": identity["email"],
        },
        "summary": selected["summary"],
        "experience": experience,
        "skills": skills,
        "projects": projects,
        "education": [
            {"degree": x["degree"], "institution": x["school"], "dates": ""}
            for x in profile["education"]
        ],
    }
    cover = None
    if package["cover"]:
        cover = {
            **package["cover"],
            "header": {
                "name": identity["name"],
                "address": identity["location"],
                "phone": identity["phone"],
                "email": identity["email"],
                "date": "",
            },
            "salutation": "Dear Hiring Manager,",
            "sign_off": "Sincerely,\n" + identity["name"],
        }
    public_text_check([resume, cover], identity["email"])
    return resume, cover


def current_display_label(function):
    """Return the approved resume-facing label for an unconfirmed current function."""
    labels = {
        "capacity and performance specialist": "Capacity & Performance Specialist",
    }
    return labels.get(function.casefold(), function)


def load_renderer(root, filename):
    spec = importlib.util.spec_from_file_location(
        "canonical_" + filename[:-3], root / "scripts" / filename
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def render_local(root: Path, directory: Path, cover_requested: bool):
    """Only pure rendering functions are called; legacy CLI/master/LLM paths are bypassed."""
    from docx import Document

    resume = read_json(directory / "resume_intermediate_v1.json")
    renderer = load_renderer(root, "05_render_resume.py")
    renderer.render_docx(
        resume, directory / "resume.docx", trim=False, exclusions=set()
    )
    preview = renderer.render_markdown(resume, trim=False, exclusions=set())
    # The old renderer assumes a start date. Preserve the explicitly unknown current start.
    doc = Document(directory / "resume.docx")
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.text = run.text.replace(" ( – Present)", " (Current)")
    doc.save(directory / "resume.docx")
    atomic_text(
        directory / "resume_preview_v1.md",
        preview.replace("\n – Present\n", "\nCurrent\n"),
    )
    if cover_requested:
        cover = read_json(directory / "cover_intermediate_v1.json")
        renderer = load_renderer(root, "08_render_cover_letter.py")
        renderer.render_docx(cover, directory / "cover.docx")
        atomic_text(directory / "cover_preview_v1.md", renderer.render_markdown(cover))
    quality = load_renderer(root, "quality_check.py")
    issues = quality.check_resume(directory, "v1")
    if cover_requested:
        issues += quality.check_cover(directory, "v1")
    for name in (["resume.docx", "cover.docx"] if cover_requested else ["resume.docx"]):
        document = Document(directory / name)
        text = "\n".join(p.text for p in document.paragraphs)
        public_text_check(text, resume["personal"]["email"])
        if resume["personal"]["full_name"] not in text or not text.strip():
            issues.append("Rendered document missing identity/content")
    result = {
        "passed": not issues,
        "issues": issues,
        "type": "local_structural_and_content",
        "visual_review_required": True,
        "note": "Open the DOCX and review every page before submitting; no visual verification is claimed.",
    }
    write_json(directory / "quality_report.json", result)
    return result


def safe_artifact_path(folder, relative):
    path = (folder / relative).resolve()
    if not path.is_relative_to((folder / "generated/packages").resolve()):
        raise PipelineError("Artifact path is outside this job's generated packages.")
    return path


def verify_artifacts(folder, artifacts):
    manifest_path = safe_artifact_path(folder, artifacts["manifest"])
    if digest(manifest_path.read_bytes()) != artifacts["manifest_sha256"]:
        raise PipelineError(
            "Artifact manifest changed; preserve edits and review before recording application."
        )
    manifest = read_json(manifest_path)
    if not manifest.get("quality", {}).get("passed") or not manifest.get("files"):
        raise PipelineError("Artifact package has not passed local quality checks.")
    for filename, expected in manifest["files"].items():
        path = (manifest_path.parent / filename).resolve()
        if path.parent != manifest_path.parent or digest(path.read_bytes()) != expected:
            raise PipelineError(
                "Generated artifacts changed or are missing; preserve edits and review."
            )
    for key in ("resume", "cover"):
        if artifacts.get(key):
            path = safe_artifact_path(folder, artifacts[key])
            if (
                path.parent != manifest_path.parent
                or path.name not in manifest["files"]
            ):
                raise PipelineError(
                    "Artifact path does not match the completed manifest."
                )
    return manifest


def prepare_artifacts(root, folder, package, profile, profile_hash, provenance):
    """Each package has a separate folder; repeated cache hits never overwrite edits."""
    renderer_sources = [
        root / "scripts" / name
        for name in (
            "05_render_resume.py",
            "08_render_cover_letter.py",
            "quality_check.py",
        )
    ]
    renderer_hash = digest(
        canonical(
            [digest(path.read_bytes()) for path in [Path(__file__), *renderer_sources]]
        )
    )
    selection_hash = digest(canonical(package))
    version = digest(
        canonical({"cache_key": provenance["tier2"]["key"], "renderers": renderer_hash})
    )
    directory = folder / "generated/packages" / version
    manifest_path = directory / "artifact_manifest.json"
    paths = {
        "resume": (directory / "resume.docx").relative_to(folder).as_posix(),
        "cover": (
            (directory / "cover.docx").relative_to(folder).as_posix()
            if package["cover"]
            else None
        ),
        "manifest": manifest_path.relative_to(folder).as_posix(),
    }
    if directory.exists():
        if not manifest_path.is_file():
            raise PipelineError(
                "An incomplete local render exists; preserve it and diagnose before retrying."
            )
        paths["manifest_sha256"] = digest(manifest_path.read_bytes())
        # If this is the active package, also verify the independently saved manifest hash.
        import yaml

        metadata = (
            yaml.safe_load((folder / "metadata.yaml").read_text(encoding="utf-8")) or {}
        )
        prior = metadata.get("artifacts", {})
        if (
            prior.get("manifest") == paths["manifest"]
            and prior.get("manifest_sha256") != paths["manifest_sha256"]
        ):
            raise PipelineError(
                "Existing artifact manifest was edited; nothing overwritten."
            )
        manifest = verify_artifacts(folder, paths)
        if (
            manifest.get("input_sha256") != selection_hash
            or manifest.get("runtime_profile_hash") != profile_hash
        ):
            raise PipelineError("Artifact package identity mismatch.")
        return paths
    directory.mkdir(parents=True)
    resume, cover = materialize(package, profile)
    write_json(directory / "resume_intermediate_v1.json", resume)
    if cover is not None:
        write_json(directory / "cover_intermediate_v1.json", cover)
    quality = render_local(root, directory, cover is not None)
    if not quality["passed"]:
        raise PipelineError(
            f"Local quality checks failed; inspect {directory / 'quality_report.json'}."
        )
    files = {
        path.name: digest(path.read_bytes())
        for path in sorted(directory.iterdir())
        if path.is_file()
    }
    write_json(
        manifest_path,
        {
            "version": 1,
            "input_sha256": selection_hash,
            "runtime_profile_hash": profile_hash,
            "renderer_sha256": renderer_hash,
            "provenance": provenance,
            "quality": quality,
            "files": files,
        },
    )
    paths["manifest_sha256"] = digest(manifest_path.read_bytes())
    verify_artifacts(folder, paths)
    return paths
