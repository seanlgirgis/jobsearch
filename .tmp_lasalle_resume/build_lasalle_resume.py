from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

source = Path(r"D:\Workarea\jobsearch\resume_output\Sean_Girgis_Toloka_FDE_Resume.docx")
target = Path(r"D:\Workarea\jobsearch\resume_output\Sean_Girgis_LaSalle_FDE_Finance_Resume.docx")

content = [
    "Sean Girgis",
    "Murphy, TX | 214-315-2190 | seanlgirgis@gmail.com | LinkedIn | GitHub",
    "Forward Deployed Engineer | Financial Services Technology | Python, C++, Data and AI Workflows",
    "Available for full-time contract or contract-to-hire engagement, remote or hybrid in Chicago.",
    "Professional Summary",
    "Forward-deployed and data engineer with enterprise financial-services experience translating stakeholder needs into reliable, high-performance data, analytics, and workflow solutions. Hands-on with Python, SQL, C++, APIs, ETL, performance engineering, and AI-assisted development; experienced working with operational teams from discovery through delivery and support.",
    "Solution Engineering",
    "Stakeholder delivery: discovery, requirements translation, solution design, client-facing technical delivery, SDLC collaboration",
    "Software and data: Python, C++, SQL, APIs, Oracle, Microsoft SQL Server, ETL, data validation, analytics workflows",
    "Financial systems and performance: capacity planning, forecasting, observability, performance analysis, scalable transaction systems",
    "Modern AI workflow exposure: LLM-assisted development, RAG, FAISS, agentic workflows, AI evaluation, searchable knowledge bases",
    "Selected Outcomes and Projects",
    "Built Python and Pandas ETL pipelines for P95 telemetry from 6,000+ enterprise endpoints, with validation, cleansing, anomaly checks, and Oracle history for capacity and risk analysis.",
    "Developed Prophet and scikit-learn forecasting models to surface capacity risks three to six months ahead and support provisioning and rightsizing decisions.",
    "Led migration of a high-throughput shopping engine from 200+ MySQL nodes to a six-node Oracle RAC cluster; optimized C++/OCCI transaction processing while maintaining sub-second latency.",
    "Built an agentic job-search workflow with Python, FAISS, sentence-transformers, LLM scoring and tailoring, document generation, quality gates, and application tracking.",
    "Automated capacity-report preparation and integrated JIRA, Microsoft SQL Server, and Splunk inputs to standardize recurring analysis and improve report quality.",
    "Professional Experience",
    "LTIMindtree | Capacity and Performance Specialist\t2026 - Present",
    "Use practical AI and automation to improve capacity reporting, integrate operational data sources, and make recurring analyst workflows more consistent and efficient.",
    "Research a specialized team knowledge-base and analyst command-center concept using LLMs, retrieval, Python, and PowerShell; this work remains investigative rather than a production deployment.",
    "Citi | Senior Capacity and Data Engineer\tNov 2017 - Dec 2025",
    "Architected Python and Pandas ETL pipelines for P95 telemetry from 6,000+ endpoints, adding validation, cleansing, anomaly checks, and durable Oracle history.",
    "Developed forecasting models to identify capacity risks three to six months ahead, supporting planning decisions in a financial-services environment.",
    "Managed capacity planning and observability across BMC TrueSight, CA APM, and AppDynamics; produced trend reporting, seasonal analysis, and executive dashboards.",
    "G6 Hospitality | Performance Engineer\tMar 2017 - Nov 2017",
    "Managed Dynatrace AppMon and Synthetics for critical digital systems, analyzed bottlenecks, supported an AWS migration, and led performance-data mining for optimization recommendations.",
    "Earlier Consulting and Engineering Experience | APM, Performance, Data Migration and Software Engineering\t1996 - 2017",
    "Delivered enterprise APM implementations and upgrades for estates of roughly 4,000-6,000 agents; created monitoring modules, dashboards, alerts, extraction scripts, sizing guidance, and client training.",
    "Built high-availability multithreaded C++ interfaces and supported performance, data migration, and production problem solving across enterprise systems.",
    "Technical Skills",
    "Python, C++, Pandas, APIs, SQL, Oracle, Microsoft SQL Server, ETL/ELT, PowerShell, Linux/Unix, AWS S3/Glue/Athena/Bedrock, PySpark, Airflow, Docker, Parquet, Splunk, Power BI, CA APM, Dynatrace, AppDynamics, BMC TrueSight, Prophet, scikit-learn, Streamlit, RAG, FAISS, MCP, solution engineering, AI evaluation, LLM and agentic workflows",
    "Education and Certifications",
    "Post-Graduate Diploma, Computer Engineering Technology / Computer Science - Humber College, Toronto\nBachelor of Science, Civil Engineering - Zagazig University, Egypt\nBrainbench certifications: C++ Fundamentals, C++, C, Java 2.0, Object-Oriented Design, Mathematics",
]

doc = Document(source)
assert len(doc.paragraphs) == len(content), (len(doc.paragraphs), len(content))
for index, (paragraph, text) in enumerate(zip(doc.paragraphs, content)):
    # The contact paragraph includes hyperlink field XML; retaining it avoids
    # duplicating the visible LinkedIn and GitHub labels.
    if index == 1:
        continue
    # Preserve paragraph styles/formatting and replace only the visible text.
    for run in paragraph.runs:
        run.text = ""
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run.text = text
    # Ensure the document retains the source's font family when a renderer opens it.
    if run.font.name:
        run._element.rPr.rFonts.set(qn('w:ascii'), run.font.name)
        run._element.rPr.rFonts.set(qn('w:hAnsi'), run.font.name)
doc.core_properties.title = "Sean Girgis Resume Forward Deployed Engineer Finance"
doc.core_properties.subject = "Resume for LaSalle Network Forward Deployed Engineer Finance and Investments"
doc.core_properties.author = "Sean Girgis"
doc.save(target)
print(target)
